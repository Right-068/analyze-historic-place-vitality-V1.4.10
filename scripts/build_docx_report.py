#!/usr/bin/env python3
"""Build the formal DOCX report from structured report data and the source ledger."""

from __future__ import annotations

import argparse
import csv
import strict_json as json
import re
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path

from dimension_evidence import (
    HIGH_EVIDENCE_UNITS,
    HIGH_SOURCE_CATEGORIES,
    HIGH_SOURCE_PAGES,
    MEDIUM_HIGH_EVIDENCE_UNITS,
    MEDIUM_HIGH_SOURCE_CATEGORIES,
    MEDIUM_HIGH_SOURCE_PAGES,
    MIN_DIMENSION_EVIDENCE_UNITS,
    MIN_DIMENSION_SOURCE_CATEGORIES,
    MIN_DIMENSION_SOURCE_PAGES,
    MIN_MEDIUM_ENHANCEMENT_ROUNDS,
)
from dimension_framework import (
    DIMENSION_BOUNDARIES,
    DIMENSION_DEFINITIONS,
    DIMENSION_NAMES,
    DIMENSION_WEIGHTS,
    RESULT_LAYER_NAME,
    WEIGHT_SCHEME_NAME,
)
from formal_states import AUDIT_RETRIEVAL_TERMINATED_WITH_SHORTFALL
from runtime_guard import verify_scoring_input_gate

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - dependency guard
    raise SystemExit(
        "python-docx is required to build the DOCX deliverable. Install it in the current code runtime."
    ) from exc

from artifact_provenance import (
    assert_frozen_artifact_path,
    register_protected_artifact,
    verify_artifact_writer,
    verify_preflight,
    verify_truth_freeze,
)
from runtime_guard import authorize_runtime_write


SECTIONS = [
    "地点基本信息",
    "检索过程与数据来源",
    "来源结构分析",
    "总体情感与综合评价",
    "七个维度的详细分析",
    "七维评价汇总表",
    "高频正面评价主题",
    "高频负面评价主题",
    "不同平台评价差异",
    "不同用户群体和场景差异",
    "时间变化分析",
    "历史文化活态传承感知评价",
    "结论",
    "来源清单",
    "机器可读结果",
]

DIMENSIONS = DIMENSION_NAMES

MENTION_LEVELS = {"A", "B", "C", "D", "E", "U"}
EVIDENCE_STRENGTHS = {"高", "中高", "中", "低", "数据不足"}
RESULT_LAYER_LABELS = ["综合分", "核心优势", "核心风险", "底线性问题", "证据充分度", "评价置信度"]
MIN_POSITIVE_CJK = 100
MIN_NEGATIVE_CJK = 90
MIN_FACT_CJK = 100
MIN_DIFFERENCE_CJK = 80
MIN_DIFFERENCE_ITEMS = 2
MIN_MECHANISM_CJK = 120
MIN_COUNTEREVIDENCE_CJK = 100
MIN_UNCERTAINTY_CJK = 100
MIN_SYNTHESIS_CJK = 120
MIN_DIMENSION_CJK = 1000
MIN_SCOPE_CJK = 300
MIN_CONCLUSION_CJK = 1500
MIN_CONCLUSION_PARAGRAPHS = 7

ENGLISH_TERM_PATTERN = re.compile(
    r"(?<![A-Za-z0-9_])([A-Za-z]{2,}(?:[A-Za-z0-9+./-]*)(?:\s+[A-Za-z][A-Za-z0-9+./-]*){0,5})(?![A-Za-z0-9_])"
)
ALLOWED_REPORT_FONTS = {"黑体", "宋体", "Times New Roman"}
DEFAULT_EVALUATION_PROTOCOL = Path(__file__).resolve().parents[1] / "assets" / "evaluation-protocol.json"

SECTION_MIN_CJK = {
    1: 800,
    2: 1000,
    3: 350,
    4: 350,
    6: 300,
    7: 250,
    8: 280,
    9: 750,
    10: 500,
    11: 650,
    12: 750,
    13: MIN_CONCLUSION_CJK,
}

REQUIRED_SUBHEADINGS = {
    1: ["扩展检索名称和关联词"],
    2: ["2.1 实际使用的检索词", "2.2 来源规模", "2.3 去重规则执行情况", "2.4 无法充分访问的平台"],
    4: ["跨平台计算"],
    6: ["权重核算"],
    9: ["9.1 平台特征", "9.2 平台评分表"],
    11: [
        "11.1 改造前后的空间变化",
        "11.2 近年用户感知的变化",
        "11.3 商业化程度变化",
        "11.4 游客数量和拥挤变化",
        "11.5 文化活动变化",
        "11.6 长期问题判断",
    ],
}

SECTION_TABLE_HEADERS = {
    1: [["项目", "基本信息"]],
    2: [
        ["指标", "本次检索结果"],
        ["平台", "访问情况", "处理方式"],
    ],
    3: [["来源平台", "来源数量", "内容类型", "时间范围", "是否用于情感分析", "主要价值", "主要局限"]],
    4: [["指标", "得分", "说明"]],
    6: [["评价维度", "初始研究权重", "正面主题", "负面主题", "提及率等级", "倾向得分", "转换分", "证据充分度", "主要来源平台"]],
    7: [["排名", "正面主题", "具体含义", "提及率等级", "正面强度1—5", "主要平台", "证据一致性"]],
    8: [["排名", "负面主题", "具体问题", "提及率等级", "负面强度1—5", "主要平台", "证据一致性"]],
    9: [["平台", "主要正面主题", "主要负面主题", "七维综合分", "样本充分度", "平台偏向"]],
    10: [["用户群体或场景", "主要感知", "主要问题", "判断可靠性"]],
    11: [["问题", "是否长期存在", "判断"]],
}

SECTION_TABLE_ROW_RANGES = {
    1: [(8, None)],
    2: [(11, None), (5, None)],
    3: [(8, None)],
    4: [(8, None)],
    6: [(7, 7)],
    7: [(5, 10)],
    8: [(5, 10)],
    9: [(3, None)],
    10: [(8, None)],
    11: [(5, None)],
}

NAVY = "17324D"
LIGHT_BLUE = "D9EAF7"
SOURCE_TABLE_HEADERS = [
    "序号",
    "平台或网站",
    "页面标题",
    "发布时间",
    "内容类型",
    "是否用于评分",
    "来源链接",
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [
            {key: (value or "").strip() for key, value in row.items()}
            for row in csv.DictReader(handle)
        ]


def set_run_font(
    run: object,
    name: str = "Times New Roman",
    size: float = 10.5,
    bold: bool | None = None,
    east_asia: str | None = None,
) -> None:
    latin_name = "Times New Roman" if name in {"黑体", "宋体"} else name
    run.font.name = latin_name
    east_asia_name = east_asia or ("宋体" if name == "Times New Roman" else name)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin_name)
    fonts.set(qn("w:hAnsi"), latin_name)
    fonts.set(qn("w:eastAsia"), east_asia_name)
    fonts.set(qn("w:cs"), latin_name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell: object, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell: object, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph: object, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.extend([color, underline, fonts, size])
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(document: object) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)
    body_text = document.styles["Body Text"]
    body_text.font.name = "Times New Roman"
    body_text._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    body_text.font.size = Pt(10.5)
    body_text.paragraph_format.line_spacing = 1.5
    body_text.paragraph_format.space_after = Pt(4)
    for style_name, font_name, east_asia, size, color in (
        ("Title", "Times New Roman", "黑体", 24, NAVY),
        ("Subtitle", "Times New Roman", "黑体", 14, "666666"),
        ("Heading 1", "Times New Roman", "黑体", 16, NAVY),
        ("Heading 2", "Times New Roman", "黑体", 14, NAVY),
        ("Heading 3", "Times New Roman", "黑体", 12, NAVY),
    ):
        style = document.styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name.startswith("Heading") or style_name == "Title"
        if style_name == "Title":
            style.paragraph_format.line_spacing = 1.0
            style.paragraph_format.space_before = Pt(70)
            style.paragraph_format.space_after = Pt(20)
        elif style_name == "Subtitle":
            style.paragraph_format.space_after = Pt(14)
        else:
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
    if "Date" not in [style.name for style in document.styles]:
        date_style = document.styles.add_style("Date", WD_STYLE_TYPE.PARAGRAPH)
    else:
        date_style = document.styles["Date"]
    date_style.font.name = "Times New Roman"
    date_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    date_style.font.size = Pt(11)
    if "TOC Heading" not in [style.name for style in document.styles]:
        toc_style = document.styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        toc_style = document.styles["TOC Heading"]
    toc_style.font.name = "Times New Roman"
    toc_style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    toc_style.font.size = Pt(16)
    toc_style.font.bold = True
    toc_style.font.color.rgb = RGBColor.from_string(NAVY)
    toc_style.paragraph_format.space_before = Pt(12)
    toc_style.paragraph_format.space_after = Pt(10)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    set_run_font(run, "宋体", 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    suffix = footer.add_run(" 页")
    set_run_font(suffix, "宋体", 9)


def add_paragraph_with_links(document: object, block: dict[str, object], style: str | None = None) -> object:
    paragraph = document.add_paragraph(style=style)
    text = str(block.get("text", ""))
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    links = block.get("links", [])
    if isinstance(links, list):
        for link in links:
            if not isinstance(link, dict):
                continue
            label = str(link.get("label", "来源")).strip() or "来源"
            url = str(link.get("url", "")).strip()
            if re.match(r"^https?://", url):
                separator = paragraph.add_run(" ")
                set_run_font(separator)
                add_hyperlink(paragraph, label, url)
    return paragraph


def add_table(document: object, headers: list[object], rows: list[list[object]]) -> object:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, "黑体", 9, bold=True)
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for column, value in enumerate(values[: len(headers)]):
            cells[column].text = "" if value is None else str(value)
            cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cells[column])
            if row_index % 2 == 0:
                set_cell_shading(cells[column], "F3F6F9")
            for paragraph in cells[column].paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, "宋体", 9)
    return table


def add_block(document: object, block: dict[str, object]) -> None:
    block_type = str(block.get("type", "paragraph"))
    if block_type == "paragraph":
        add_paragraph_with_links(document, block)
    elif block_type == "subheading":
        text = str(block.get("text", ""))
        paragraph = document.add_heading(text, level=2)
        for run in paragraph.runs:
            set_run_font(run, "Times New Roman", 14, bold=True, east_asia="黑体")
    elif block_type == "subheading3":
        text = str(block.get("text", ""))
        paragraph = document.add_heading(text, level=3)
        for run in paragraph.runs:
            set_run_font(run, "Times New Roman", 12, bold=True, east_asia="黑体")
    elif block_type in {"bullets", "numbered"}:
        items = block.get("items", [])
        if not isinstance(items, list):
            raise ValueError(f"{block_type} block requires an items array")
        style = "List Bullet" if block_type == "bullets" else "List Number"
        for item in items:
            payload = item if isinstance(item, dict) else {"text": str(item)}
            add_paragraph_with_links(document, payload, style=style)
    elif block_type == "quote":
        paragraph = add_paragraph_with_links(document, block)
        paragraph.paragraph_format.left_indent = Cm(0.8)
        paragraph.paragraph_format.right_indent = Cm(0.8)
        set_cell_like_border(paragraph)
    elif block_type == "table":
        headers = block.get("headers", [])
        rows = block.get("rows", [])
        if not isinstance(headers, list) or not headers:
            raise ValueError("table block requires a non-empty headers array")
        if not isinstance(rows, list):
            raise ValueError("table block requires a rows array")
        add_table(document, headers, [row for row in rows if isinstance(row, list)])
    elif block_type == "key_values":
        items = block.get("items", [])
        if not isinstance(items, list):
            raise ValueError("key_values block requires an items array")
        rows = [[item.get("label", ""), item.get("value", "")] for item in items if isinstance(item, dict)]
        add_table(document, ["项目", "内容"], rows)
    elif block_type == "judgments":
        items = block.get("items", [])
        if not isinstance(items, list):
            raise ValueError("judgments block requires an items array")
        for item in items:
            if not isinstance(item, dict):
                raise ValueError("each judgments item must be an object")
            label = str(item.get("label", "")).strip()
            value = str(item.get("value", "")).strip()
            paragraph = document.add_paragraph()
            label_run = paragraph.add_run(f"{label}：")
            set_run_font(label_run, "宋体", 10.5, bold=True)
            value_run = paragraph.add_run(value)
            set_run_font(value_run, "宋体", 10.5)
            links = item.get("links", [])
            if isinstance(links, list):
                for link in links:
                    if not isinstance(link, dict):
                        continue
                    link_label = str(link.get("label", "来源")).strip() or "来源"
                    url = str(link.get("url", "")).strip()
                    if re.match(r"^https?://", url):
                        spacer = paragraph.add_run(" ")
                        set_run_font(spacer)
                        add_hyperlink(paragraph, link_label, url)
    elif block_type == "page_break":
        document.add_page_break()
    else:
        raise ValueError(f"unknown report block type: {block_type}")


def set_cell_like_border(paragraph: object) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), NAVY)
    borders.append(left)
    paragraph_properties.append(borders)


def add_source_table(document: object, sources: list[dict[str, str]]) -> None:
    intro = document.add_paragraph(
        f"以下列出来源台账中的全部 {len(sources)} 个页面记录；页面编号、读取状态、检索记录和编码关系须结合 XLSX 工作簿复核。"
    )
    for run in intro.runs:
        set_run_font(run)
    table = document.add_table(rows=1, cols=len(SOURCE_TABLE_HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(SOURCE_TABLE_HEADERS):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, "黑体", 8.5, bold=True)
                run.font.color.rgb = RGBColor(255, 255, 255)
    for index, source in enumerate(sources, start=1):
        cells = table.add_row().cells
        values = [
            index,
            source.get("platform", ""),
            source.get("page_title", ""),
            source.get("published_at", ""),
            source.get("source_category", ""),
            source.get("used_for_scoring", ""),
        ]
        for column, value in enumerate(values):
            cells[column].text = str(value)
        url = source.get("url", "")
        cells[6].text = ""
        if re.match(r"^https?://", url):
            add_hyperlink(cells[6].paragraphs[0], "打开来源", url)
        else:
            cells[6].text = url
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, top=55, bottom=55)
            if index % 2 == 0:
                set_cell_shading(cell, "F3F6F9")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, "宋体", 8.5)


def cjk_count(text: str) -> int:
    return len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff]", text))


def normalize_narrative(text: str) -> str:
    without_urls = re.sub(r"https?://\S+", "", text, flags=re.IGNORECASE)
    return "".join(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fffA-Za-z]", without_urls)).casefold()


def character_ngrams(text: str, size: int = 5) -> set[str]:
    return {text[index : index + size] for index in range(max(0, len(text) - size + 1))}


def narrative_similarity(left: str, right: str) -> tuple[float, float]:
    a = normalize_narrative(left)
    b = normalize_narrative(right)
    if not a or not b:
        return 0.0, 0.0
    ratio = SequenceMatcher(None, a, b, autojunk=False).ratio()
    a_grams = character_ngrams(a)
    b_grams = character_ngrams(b)
    union = a_grams | b_grams
    jaccard = len(a_grams & b_grams) / len(union) if union else 0.0
    return ratio, jaccard


def assert_no_near_duplicate_narratives(
    records: list[tuple[str, str]],
    *,
    context: str,
) -> None:
    eligible = [
        (label, text)
        for label, text in records
        if cjk_count(text) >= 60 and len(normalize_narrative(text)) >= 80
    ]
    for index, (left_label, left) in enumerate(eligible):
        left_normalized = normalize_narrative(left)
        for right_label, right in eligible[index + 1 :]:
            right_normalized = normalize_narrative(right)
            if left_normalized == right_normalized:
                raise ValueError(
                    f"{context} contains duplicated prose between {left_label} and {right_label}"
                )
            ratio, jaccard = narrative_similarity(left, right)
            if ratio >= 0.92 or (ratio >= 0.86 and jaccard >= 0.78):
                raise ValueError(
                    f"{context} contains near-duplicate templated prose between {left_label} and {right_label}; "
                    f"similarity={ratio:.3f}, ngram_jaccard={jaccard:.3f}"
                )


def english_gloss_violations(text: str) -> list[str]:
    cleaned = re.sub(r"https?://\S+|\b[\w.-]+@[\w.-]+\b", " ", text)
    cleaned = re.sub(
        r"任务运行编号(?:为|[:：])?\s*[A-Za-z][A-Za-z0-9._-]*",
        "任务运行编号",
        cleaned,
    )
    cleaned = re.sub(
        r"(?<![A-Za-z0-9])(?=[A-Za-z0-9._-]*\d)[A-Za-z][A-Za-z0-9._-]*(?![A-Za-z0-9])",
        " ",
        cleaned,
    )
    cleaned = re.sub(r"\b[A-Za-z]+_[A-Za-z0-9_-]+\b", " ", cleaned)
    violations: list[str] = []
    for match in ENGLISH_TERM_PATTERN.finditer(cleaned):
        term = match.group(1).strip()
        if re.fullmatch(r"[A-E](?:\s*[-—]\s*[A-E])?", term, flags=re.IGNORECASE):
            continue
        after = cleaned[match.end() : match.end() + 80]
        if re.match(r"\s*[（(][^）)]*[\u3400-\u4dbf\u4e00-\u9fff][^）)]*[）)]", after):
            continue
        before = cleaned[: match.start()]
        open_index = max(before.rfind("（"), before.rfind("("))
        close_index = max(before.rfind("）"), before.rfind(")"))
        if open_index > close_index:
            preceding = before[max(0, open_index - 30) : open_index]
            closing = cleaned[match.end() : match.end() + 80]
            if cjk_count(preceding) > 0 and re.search(r"[）)]", closing):
                continue
        snippet = cleaned[max(0, match.start() - 18) : min(len(cleaned), match.end() + 18)]
        violations.append(f"{term}（上下文：{snippet}）")
    return violations


def validate_english_gloss_contract(data: dict[str, object]) -> None:
    prose: list[str] = []
    scope = data.get("research_scope", [])
    if isinstance(scope, list):
        prose.extend(text_payload(item) for item in scope)
    sections = data.get("sections", [])
    if isinstance(sections, list):
        for section in sections:
            if not isinstance(section, dict) or not isinstance(section.get("number"), int):
                continue
            if int(section["number"]) > 13:
                continue
            blocks = section.get("blocks", [])
            if isinstance(blocks, list):
                prose.extend(block_text(block) for block in blocks if isinstance(block, dict))
    dimensions = data.get("dimension_analyses", [])
    if isinstance(dimensions, list):
        for item in dimensions:
            if not isinstance(item, dict):
                continue
            for field in (
                "positive", "negative", "differences", "fact_perception", "mechanism_analysis",
                "counterevidence", "uncertainty_boundary", "synthesis_judgment",
            ):
                values = item.get(field, [])
                if isinstance(values, list):
                    prose.extend(text_payload(value) for value in values)
            prose.extend(
                str(item.get(field, ""))
                for field in ("differences_title", "mention_range", "mention_basis")
            )
            additional = item.get("additional_analysis", [])
            if isinstance(additional, list):
                prose.extend(block_text(block) for block in additional if isinstance(block, dict))
    violations = english_gloss_violations("\n".join(prose))
    if violations:
        preview = "；".join(violations[:8])
        raise ValueError(
            "report body contains English without an adjacent Chinese gloss; use English（中文释义） "
            f"or 中文释义（English）: {preview}"
        )


def text_payload(value: object) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        return str(value.get("text", ""))
    return ""


def block_text(block: dict[str, object]) -> str:
    block_type = str(block.get("type", "paragraph"))
    if block_type in {"paragraph", "subheading", "subheading3", "quote"}:
        return str(block.get("text", ""))
    if block_type in {"bullets", "numbered"}:
        items = block.get("items", [])
        return "".join(text_payload(item) for item in items) if isinstance(items, list) else ""
    if block_type == "table":
        headers = block.get("headers", [])
        rows = block.get("rows", [])
        parts = [str(value) for value in headers if value is not None] if isinstance(headers, list) else []
        if isinstance(rows, list):
            for row in rows:
                if isinstance(row, list):
                    parts.extend(str(value) for value in row if value is not None)
        return "".join(parts)
    if block_type == "key_values":
        items = block.get("items", [])
        if not isinstance(items, list):
            return ""
        return "".join(
            f"{item.get('label', '')}{item.get('value', '')}"
            for item in items
            if isinstance(item, dict)
        )
    if block_type == "judgments":
        items = block.get("items", [])
        if not isinstance(items, list):
            return ""
        return "".join(
            f"{item.get('label', '')}{item.get('value', '')}"
            for item in items
            if isinstance(item, dict)
        )
    return ""


def validate_research_scope(data: dict[str, object]) -> list[object]:
    scope = data.get("research_scope")
    if not isinstance(scope, list) or not 4 <= len(scope) <= 6:
        raise ValueError("research_scope must contain four to six evidence-boundary paragraphs")
    for index, item in enumerate(scope, start=1):
        text = text_payload(item).strip()
        if not text:
            raise ValueError(f"research_scope item {index} text is required")
        if isinstance(item, dict) and not isinstance(item.get("links", []), list):
            raise ValueError(f"research_scope item {index} links must be an array")
    if cjk_count("".join(text_payload(item) for item in scope)) < MIN_SCOPE_CJK:
        raise ValueError(f"research_scope must contain at least {MIN_SCOPE_CJK} Chinese characters")
    return scope


def validate_section_contract(number: int, blocks: list[object]) -> None:
    typed = [block for block in blocks if isinstance(block, dict)]
    if len(typed) != len(blocks):
        raise ValueError(f"section {number} blocks must contain objects only")
    tables = [block for block in typed if block.get("type") == "table"]
    expected_headers = SECTION_TABLE_HEADERS.get(number, [])
    actual_headers = [block.get("headers", []) for block in tables]
    if actual_headers != expected_headers:
        raise ValueError(
            f"section {number} table headers must be exactly {expected_headers}; found {actual_headers}"
        )
    ranges = SECTION_TABLE_ROW_RANGES.get(number, [])
    for table_index, (table, limits) in enumerate(zip(tables, ranges), start=1):
        rows = table.get("rows", [])
        if not isinstance(rows, list):
            raise ValueError(f"section {number} table {table_index} rows must be an array")
        minimum, maximum = limits
        if len(rows) < minimum or (maximum is not None and len(rows) > maximum):
            upper = "unbounded" if maximum is None else str(maximum)
            raise ValueError(
                f"section {number} table {table_index} must contain {minimum} to {upper} data rows"
            )
    actual_subheadings = [
        str(block.get("text", "")).strip()
        for block in typed
        if block.get("type") == "subheading"
    ]
    expected_subheadings = REQUIRED_SUBHEADINGS.get(number, [])
    if actual_subheadings != expected_subheadings:
        raise ValueError(
            f"section {number} Heading 2 sequence must be exactly {expected_subheadings}; found {actual_subheadings}"
        )
    if number == 9:
        feature_index = next(
            index for index, block in enumerate(typed)
            if block.get("type") == "subheading" and block.get("text") == "9.1 平台特征"
        )
        table_index = next(
            index for index, block in enumerate(typed)
            if block.get("type") == "subheading" and block.get("text") == "9.2 平台评分表"
        )
        platform_headings = [
            block for block in typed[feature_index + 1 : table_index]
            if block.get("type") == "subheading3" and str(block.get("text", "")).strip()
        ]
        if len(platform_headings) < 3:
            raise ValueError("section 9 must contain at least three evidence-based platform or source-group Heading 3 blocks")
    elif any(block.get("type") == "subheading3" for block in typed):
        raise ValueError(f"section {number} may not contain Heading 3 blocks")
    if number == 12:
        judgment_blocks = [block for block in typed if block.get("type") == "judgments"]
        if len(judgment_blocks) != 1:
            raise ValueError("section 12 must contain exactly one judgments block")
        items = judgment_blocks[0].get("items", [])
        if not isinstance(items, list) or len(items) != len(RESULT_LAYER_LABELS):
            raise ValueError("section 12 judgments block must contain exactly the six +1 result fields")
        labels = [str(item.get("label", "")).strip() for item in items if isinstance(item, dict)]
        if labels != RESULT_LAYER_LABELS:
            raise ValueError(f"section 12 judgment labels must be {RESULT_LAYER_LABELS}")
        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict) or not str(item.get("label", "")).strip() or not str(item.get("value", "")).strip():
                raise ValueError(f"section 12 judgment {index} requires a label and substantive value")
    minimum = SECTION_MIN_CJK.get(number)
    section_cjk = cjk_count("".join(block_text(block) for block in typed))
    if minimum is not None and section_cjk < minimum:
        raise ValueError(
            f"section {number} contains only {section_cjk} Chinese characters across narrative and analytical tables; at least {minimum} required"
        )
    if number == 13:
        paragraphs = [block for block in typed if block.get("type") == "paragraph" and str(block.get("text", "")).strip()]
        if len(paragraphs) < MIN_CONCLUSION_PARAGRAPHS:
            raise ValueError(
                f"section 13 must contain at least {MIN_CONCLUSION_PARAGRAPHS} substantive conclusion paragraphs"
            )


def validate_text_items(
    value: object,
    label: str,
    *,
    minimum_items: int = 1,
    minimum_cjk: int = 1,
) -> list[object]:
    if not isinstance(value, list) or len(value) < minimum_items:
        raise ValueError(f"{label} must contain at least {minimum_items} item(s)")
    combined_text: list[str] = []
    for index, item in enumerate(value, start=1):
        if isinstance(item, str):
            text = item.strip()
        elif isinstance(item, dict):
            text = str(item.get("text", "")).strip()
            links = item.get("links", [])
            if not isinstance(links, list):
                raise ValueError(f"{label} item {index} links must be an array")
        else:
            raise ValueError(f"{label} item {index} must be text or an object")
        if not text:
            raise ValueError(f"{label} item {index} text is required")
        combined_text.append(text)
    if cjk_count("".join(combined_text)) < minimum_cjk:
        raise ValueError(f"{label} must contain at least {minimum_cjk} Chinese characters")
    return value


def text_items_text(items: list[object]) -> str:
    return "".join(
        item if isinstance(item, str) else str(item.get("text", ""))
        for item in items
        if isinstance(item, (str, dict))
    )


def require_reasoning_markers(items: list[object], label: str, markers: tuple[str, ...]) -> None:
    text = text_items_text(items)
    if not any(marker in text for marker in markers):
        raise ValueError(
            f"{label} must express an analytical relationship rather than descriptive padding; "
            f"include at least one applicable reasoning cue from {markers}"
        )


def validate_dimension_analyses(data: dict[str, object]) -> list[dict[str, object]]:
    raw = data.get("dimension_analyses")
    if not isinstance(raw, list) or len(raw) != len(DIMENSIONS):
        raise ValueError("dimension_analyses must contain exactly seven canonical dimension objects")
    normalized: list[dict[str, object]] = []
    substantive_narratives: list[tuple[str, str]] = []
    for number, expected_name in enumerate(DIMENSIONS, start=1):
        item = raw[number - 1]
        if not isinstance(item, dict):
            raise ValueError(f"dimension {number} must be an object")
        if item.get("number") != number or str(item.get("name", "")).strip() != expected_name:
            raise ValueError(f"dimension {number} must be named {expected_name}")
        if str(item.get("definition", "")).strip() != DIMENSION_DEFINITIONS[expected_name]:
            raise ValueError(f"dimension {number} definition must match the canonical framework")
        if str(item.get("boundary", "")).strip() != DIMENSION_BOUNDARIES[expected_name]:
            raise ValueError(f"dimension {number} boundary must match the canonical framework")
        initial_weight = item.get("initial_research_weight")
        if (
            isinstance(initial_weight, bool)
            or not isinstance(initial_weight, (int, float))
            or abs(float(initial_weight) - DIMENSION_WEIGHTS[expected_name]) > 1e-12
        ):
            raise ValueError(
                f"dimension {number} initial_research_weight must equal the canonical value {DIMENSION_WEIGHTS[expected_name]:.0%}"
            )
        positive = validate_text_items(
            item.get("positive"),
            f"dimension {number} positive",
            minimum_cjk=MIN_POSITIVE_CJK,
        )
        negative = validate_text_items(
            item.get("negative"),
            f"dimension {number} negative",
            minimum_cjk=MIN_NEGATIVE_CJK,
        )
        differences = validate_text_items(
            item.get("differences"),
            f"dimension {number} differences",
            minimum_items=MIN_DIFFERENCE_ITEMS,
            minimum_cjk=MIN_DIFFERENCE_CJK,
        )
        fact_perception = validate_text_items(
            item.get("fact_perception"),
            f"dimension {number} fact_perception",
            minimum_cjk=MIN_FACT_CJK,
        )
        mechanism_analysis = validate_text_items(
            item.get("mechanism_analysis"),
            f"dimension {number} mechanism_analysis",
            minimum_cjk=MIN_MECHANISM_CJK,
        )
        counterevidence = validate_text_items(
            item.get("counterevidence"),
            f"dimension {number} counterevidence",
            minimum_cjk=MIN_COUNTEREVIDENCE_CJK,
        )
        uncertainty_boundary = validate_text_items(
            item.get("uncertainty_boundary"),
            f"dimension {number} uncertainty_boundary",
            minimum_cjk=MIN_UNCERTAINTY_CJK,
        )
        synthesis_judgment = validate_text_items(
            item.get("synthesis_judgment"),
            f"dimension {number} synthesis_judgment",
            minimum_cjk=MIN_SYNTHESIS_CJK,
        )
        require_reasoning_markers(
            mechanism_analysis,
            f"dimension {number} mechanism_analysis",
            ("因为", "由于", "通过", "使得", "导致", "从而", "机制", "路径"),
        )
        require_reasoning_markers(
            counterevidence,
            f"dimension {number} counterevidence",
            ("但", "然而", "反例", "替代解释", "也可能", "不能排除", "相反"),
        )
        require_reasoning_markers(
            uncertainty_boundary,
            f"dimension {number} uncertainty_boundary",
            ("样本", "来源", "平台", "时间", "边界", "仅限", "不能外推", "不确定"),
        )
        require_reasoning_markers(
            synthesis_judgment,
            f"dimension {number} synthesis_judgment",
            ("因此", "综合", "据此", "判断", "表明", "结论"),
        )
        narrative_groups = [
            positive,
            negative,
            differences,
            fact_perception,
            mechanism_analysis,
            counterevidence,
            uncertainty_boundary,
            synthesis_judgment,
        ]
        dimension_narrative = "".join(text_items_text(group) for group in narrative_groups)
        if cjk_count(dimension_narrative) < MIN_DIMENSION_CJK:
            raise ValueError(
                f"dimension {number} detailed narrative must contain at least {MIN_DIMENSION_CJK} Chinese characters"
            )
        for group_name, group in zip(
            (
                "positive", "negative", "differences", "fact_perception", "mechanism_analysis",
                "counterevidence", "uncertainty_boundary", "synthesis_judgment",
            ),
            narrative_groups,
        ):
            substantive_narratives.append(
                (f"dimension {number} {group_name}", text_items_text(group))
            )
        differences_title = str(item.get("differences_title", "")).strip()
        if not differences_title or cjk_count(differences_title) > 14:
            raise ValueError(
                f"dimension {number} differences_title must be a concise evidence-based Heading 3"
            )
        themes = item.get("themes")
        if (
            not isinstance(themes, list)
            or not 3 <= len(themes) <= 8
            or any(not str(value).strip() for value in themes)
        ):
            raise ValueError(f"dimension {number} themes must contain three to eight text items")
        mention_level = str(item.get("mention_level", "")).strip()
        if mention_level not in MENTION_LEVELS:
            raise ValueError(f"dimension {number} mention_level must be A-E or U")
        for field in ("mention_range", "mention_basis", "evidence_strength"):
            if not str(item.get(field, "")).strip():
                raise ValueError(f"dimension {number} {field} is required")
        evidence_strength = str(item.get("evidence_strength", "")).strip()
        if evidence_strength not in EVIDENCE_STRENGTHS:
            raise ValueError(
                f"dimension {number} evidence_strength must be one of {sorted(EVIDENCE_STRENGTHS)}"
            )
        confidence = str(item.get("dimension_confidence", "")).strip()
        if confidence != evidence_strength:
            raise ValueError(
                f"dimension {number} dimension_confidence must equal evidence_strength"
            )
        if str(item.get("minimum_confidence", "")).strip() != "中":
            raise ValueError(f"dimension {number} minimum_confidence must be 中")
        retrieved_units = item.get("retrieved_evidence_units")
        topic_units = item.get("topic_coverage_evidence_units")
        eligible_units = item.get("eligible_evidence_units")
        scored_units = item.get("scored_evidence_units")
        evidence_units = item.get("evidence_units")
        positive_units = item.get("scoring_positive_units")
        neutral_units = item.get("scoring_neutral_units")
        negative_units = item.get("scoring_negative_units")
        valid_platforms = item.get("valid_scoring_platforms")
        candidate_platforms = item.get("dimension_candidate_platform_count")
        scorable_platforms = item.get("dimension_scorable_platform_count")
        run_platforms = item.get("run_included_platform_count")
        source_pages = item.get("distinct_source_pages")
        source_categories = item.get("source_category_count")
        for field, value in (
            ("retrieved_evidence_units", retrieved_units),
            ("topic_coverage_evidence_units", topic_units),
            ("eligible_evidence_units", eligible_units),
            ("scored_evidence_units", scored_units),
            ("evidence_units", evidence_units),
            ("scoring_positive_units", positive_units),
            ("scoring_neutral_units", neutral_units),
            ("scoring_negative_units", negative_units),
            ("valid_scoring_platforms", valid_platforms),
            ("dimension_candidate_platform_count", candidate_platforms),
            ("dimension_scorable_platform_count", scorable_platforms),
            ("run_included_platform_count", run_platforms),
            ("distinct_source_pages", source_pages),
            ("source_category_count", source_categories),
        ):
            if isinstance(value, bool) or not isinstance(value, int) or value < 0:
                raise ValueError(f"dimension {number} {field} must be a non-negative integer")
        if evidence_units != scored_units:
            raise ValueError(f"dimension {number} evidence_units compatibility alias must equal scored_evidence_units")
        if positive_units + neutral_units + negative_units != scored_units:
            raise ValueError(f"dimension {number} formal direction counts must sum to scored_evidence_units")
        if valid_platforms != scorable_platforms:
            raise ValueError(
                f"dimension {number} valid_scoring_platforms must equal dimension_scorable_platform_count"
            )
        if scorable_platforms > run_platforms or scorable_platforms > candidate_platforms:
            raise ValueError(f"dimension {number} platform count semantics are inconsistent")
        if str(item.get("platform_minimum_rule_status", "")) not in {
            "scored", "insufficient_platform_samples"
        }:
            raise ValueError(f"dimension {number} requires a platform_minimum_rule_status")
        evidence_status = str(item.get("evidence_status", "")).strip()
        deep_rounds = item.get("targeted_deep_rounds")
        if (
            not isinstance(deep_rounds, list)
            or any(isinstance(value, bool) or not isinstance(value, int) or value < 1 for value in deep_rounds)
        ):
            raise ValueError(f"dimension {number} targeted_deep_rounds must be an integer array")
        targeted_round_count = item.get("targeted_round_count")
        if (
            isinstance(targeted_round_count, bool)
            or not isinstance(targeted_round_count, int)
            or targeted_round_count != len(deep_rounds)
        ):
            raise ValueError(
                f"dimension {number} targeted_round_count must equal len(targeted_deep_rounds)"
            )
        independent_exhaustion = item.get("independent_exhaustion_audit_passed")
        medium_completion_audit = item.get("medium_completion_audit_passed")
        if evidence_status in {"sufficient", "sufficient_at_requested_target"}:
            if (
                scored_units < MIN_DIMENSION_EVIDENCE_UNITS
                or source_pages < MIN_DIMENSION_SOURCE_PAGES
                or source_categories < MIN_DIMENSION_SOURCE_CATEGORIES
                or valid_platforms < 1
            ):
                raise ValueError(
                    f"dimension {number} is marked sufficient below the medium-confidence threshold"
                )
            expected_confidence = (
                "高" if scored_units >= HIGH_EVIDENCE_UNITS and source_pages >= HIGH_SOURCE_PAGES and source_categories >= HIGH_SOURCE_CATEGORIES
                else "中高" if scored_units >= MEDIUM_HIGH_EVIDENCE_UNITS and source_pages >= MEDIUM_HIGH_SOURCE_PAGES and source_categories >= MEDIUM_HIGH_SOURCE_CATEGORIES
                else "中"
            )
            if confidence != expected_confidence:
                raise ValueError(
                    f"dimension {number} confidence must be {expected_confidence} for its evidence coverage"
                )
            if confidence == "中" and evidence_status == "sufficient":
                raise ValueError(
                    f"dimension {number} at medium confidence requires sufficient_at_medium_after_audit, not sufficient"
                )
            if evidence_status == "sufficient_at_requested_target" and (
                confidence not in {"中", "中高", "高"}
                or str(item.get("target_confidence", "")) != "中"
            ):
                raise ValueError(
                    f"dimension {number} requested-target completion requires an explicit medium target"
                )
            if independent_exhaustion not in {False, None}:
                raise ValueError(
                    f"dimension {number} sufficient evidence must not claim exhaustion"
                )
        elif evidence_status == "sufficient_at_medium_after_audit":
            if (
                scored_units < MIN_DIMENSION_EVIDENCE_UNITS
                or source_pages < MIN_DIMENSION_SOURCE_PAGES
                or source_categories < MIN_DIMENSION_SOURCE_CATEGORIES
                or valid_platforms < 1
                or confidence != "中"
            ):
                raise ValueError(
                    f"dimension {number} medium terminal status requires the medium threshold and medium confidence"
                )
            medium_rounds = item.get("medium_enhancement_rounds")
            if (
                not isinstance(medium_rounds, list)
                or len(set(medium_rounds)) < MIN_MEDIUM_ENHANCEMENT_ROUNDS
                or medium_completion_audit is not True
                or not str(item.get("medium_completion_basis", "")).strip()
            ):
                raise ValueError(
                    f"dimension {number} medium-confidence completion lacks the stronger independent enhancement audit"
                )
            if independent_exhaustion not in {False, None}:
                raise ValueError(f"dimension {number} medium completion must not claim low-confidence exhaustion")
        elif evidence_status == "exhausted_with_shortfall":
            if confidence not in {"低", "数据不足"}:
                raise ValueError(
                    f"dimension {number} exhausted shortfall must be low confidence or data insufficient"
                )
            if len(set(deep_rounds)) < 4:
                raise ValueError(
                    f"dimension {number} exhausted shortfall requires at least four targeted deep-search rounds"
                )
            if independent_exhaustion is not True:
                raise ValueError(
                    f"dimension {number} exhausted shortfall requires an independent audit pass"
                )
            if not str(item.get("exhaustion_basis", "")).strip():
                raise ValueError(
                    f"dimension {number} exhausted shortfall requires a machine-audited exhaustion basis"
                )
        elif evidence_status == AUDIT_RETRIEVAL_TERMINATED_WITH_SHORTFALL:
            if not str(item.get("retrieval_termination_status", "")).strip():
                raise ValueError(
                    f"dimension {number} retrieval-terminal shortfall requires a machine stop reason"
                )
            numeric_score_permitted = item.get("numeric_score_permitted")
            if not isinstance(numeric_score_permitted, bool):
                raise ValueError(
                    f"dimension {number} retrieval-terminal shortfall requires numeric_score_permitted"
                )
            if numeric_score_permitted and (
                scored_units < MIN_DIMENSION_EVIDENCE_UNITS
                or source_pages < MIN_DIMENSION_SOURCE_PAGES
                or source_categories < MIN_DIMENSION_SOURCE_CATEGORIES
                or valid_platforms < 1
                or confidence not in {"中", "中高", "高"}
            ):
                raise ValueError(
                    f"dimension {number} retrieval-terminal numeric result is below the released minimum"
                )
        else:
            raise ValueError(
                f"dimension {number} evidence_status is not a released final state; "
                "needs_iteration dimensions may not enter the report"
            )
        tendency = item.get("tendency_score")
        conversion = item.get("conversion_score")
        if tendency is None:
            if conversion is not None:
                raise ValueError(f"dimension {number} conversion_score must be null when tendency_score is null")
        else:
            if isinstance(tendency, bool) or not isinstance(tendency, (int, float)) or not -5 <= tendency <= 5:
                raise ValueError(f"dimension {number} tendency_score must be numeric from -5 to 5 or null")
            if isinstance(conversion, bool) or not isinstance(conversion, (int, float)) or not 0 <= conversion <= 100:
                raise ValueError(f"dimension {number} conversion_score must be from 0 to 100")
            expected_conversion = (tendency + 5) * 10
            if abs(float(conversion) - expected_conversion) > 1e-9:
                raise ValueError(
                    f"dimension {number} conversion_score must equal (tendency_score + 5) * 10"
                )
        numeric_result_required = evidence_status in {
            "sufficient",
            "sufficient_at_requested_target",
            "sufficient_at_medium_after_audit",
        } or (
            evidence_status == AUDIT_RETRIEVAL_TERMINATED_WITH_SHORTFALL
            and item.get("numeric_score_permitted") is True
        )
        if numeric_result_required and tendency is None:
            raise ValueError(f"dimension {number} is scoring-sufficient but its formal score is empty")
        if (
            evidence_status == "exhausted_with_shortfall"
            or (
                evidence_status == AUDIT_RETRIEVAL_TERMINATED_WITH_SHORTFALL
                and item.get("numeric_score_permitted") is False
            )
        ) and tendency is not None:
            raise ValueError(f"dimension {number} is scoring-insufficient but contains a formal score")
        sources = item.get("representative_sources")
        if not isinstance(sources, list) or not sources:
            raise ValueError(f"dimension {number} representative_sources must be a non-empty array")
        for source_index, source in enumerate(sources, start=1):
            if not isinstance(source, dict) or not str(source.get("label", "")).strip():
                raise ValueError(
                    f"dimension {number} representative source {source_index} requires a label"
                )
            url = str(source.get("url", "")).strip()
            if not re.match(r"^https?://", url):
                raise ValueError(
                    f"dimension {number} representative source {source_index} requires an absolute http(s) URL"
                )
        additional = item.get("additional_analysis", [])
        if not isinstance(additional, list):
            raise ValueError(f"dimension {number} additional_analysis must be an array")
        if any(
            not isinstance(block, dict)
            or block.get("type") in {"subheading", "subheading3", "table", "key_values", "judgments", "page_break"}
            for block in additional
        ):
            raise ValueError(
                f"dimension {number} additional_analysis may contain narrative blocks only and may not add headings or tables"
            )
        normalized.append(item)
    assert_no_near_duplicate_narratives(
        substantive_narratives,
        context="dimension narratives",
    )
    return normalized


def validate_research_run(data: dict[str, object]) -> dict[str, object]:
    run = data.get("research_run")
    if not isinstance(run, dict):
        raise ValueError("research_run must be an object")
    task_run_id = str(run.get("task_run_id", "")).strip()
    if not task_run_id:
        raise ValueError("research_run task_run_id is required")
    target = run.get("minimum_effective_samples")
    actual = run.get("actual_effective_samples")
    rounds = run.get("deep_iteration_rounds")
    status = str(run.get("sample_target_status", "")).strip()
    if isinstance(target, bool) or not isinstance(target, int) or target < 100:
        raise ValueError("research_run minimum_effective_samples must be an integer of at least 100")
    if isinstance(actual, bool) or not isinstance(actual, int) or actual < 0:
        raise ValueError("research_run actual_effective_samples must be a non-negative integer")
    eligible = run.get("eligible_evidence_units")
    scored = run.get("scored_evidence_units")
    for field, value in (("eligible_evidence_units", eligible), ("scored_evidence_units", scored)):
        if isinstance(value, bool) or not isinstance(value, int) or value < 0:
            raise ValueError(f"research_run {field} must be a non-negative integer")
    if scored != actual or eligible < scored:
        raise ValueError(
            "research_run actual_effective_samples must equal scored_evidence_units and cannot exceed eligible_evidence_units"
        )
    if isinstance(rounds, bool) or not isinstance(rounds, int) or rounds < 0:
        raise ValueError("research_run deep_iteration_rounds must be a non-negative integer")
    if actual >= target:
        if status != "target_met":
            raise ValueError("research_run sample_target_status must be target_met when the target is reached")
    elif status not in {
        "exhausted_with_shortfall",
        "retrieval_terminated_with_shortfall",
    } or (status == "exhausted_with_shortfall" and rounds < 3):
            raise ValueError(
                "sample shortfall requires at least three deep rounds with machine-audited "
                "exhaustion, or a bounded retrieval-terminal state"
            )
    semantic = run.get("semantic_quantification")
    if not isinstance(semantic, dict):
        raise ValueError("research_run semantic_quantification is required")
    protocol = json.loads(DEFAULT_EVALUATION_PROTOCOL.read_text(encoding="utf-8-sig"))
    if not isinstance(protocol, dict):
        raise ValueError("released evaluation protocol root must be an object")
    if run.get("formal_scoring_schema_version") != protocol.get("formal_scoring", {}).get("schema_version"):
        raise ValueError("research_run formal_scoring_schema_version must match the released protocol")
    for field in ("evaluation_protocol_version", "semantic_codebook_version", "semantic_codebook_sha256"):
        if str(semantic.get(field, "")) != str(protocol.get(field, "")):
            raise ValueError(f"research_run semantic_quantification {field} must match the released protocol")
    if str(semantic.get("protocol_released_at", "")) != str(protocol.get("released_at", "")):
        raise ValueError("research_run semantic_quantification protocol_released_at must match the released protocol")
    if semantic.get("task_rule_refresh_performed") is not False:
        raise ValueError("research_run must state that no task-level rule refresh was performed")
    if semantic.get("task_rule_mutation_performed") is not False:
        raise ValueError("research_run must state that no task-level rule mutation was performed")
    if not str(semantic.get("semantic_rules_workbook", "")).strip():
        raise ValueError("research_run semantic_quantification semantic_rules_workbook is required")
    count_fields = (
        "native_numeric_records", "rule_codebook_records", "coding_parse_records",
        "lexicon_zero_hit_records", "coding_candidate_records", "coding_unresolved_records",
        "assisted_or_manual_records", "auto_eligible_records", "human_confirmed_records",
        "review_required_records", "low_confidence_records", "figurative_risk_records",
    )
    for field in count_fields:
        value = semantic.get(field)
        if isinstance(value, bool) or not isinstance(value, int) or value < 0:
            raise ValueError(f"research_run semantic_quantification {field} must be a non-negative integer")
    return run


def validate_representative_sources_in_ledger(
    analyses: list[dict[str, object]], sources: list[dict[str, str]]
) -> None:
    ledger_urls = {str(source.get("url", "")).strip() for source in sources}
    for number, item in enumerate(analyses, start=1):
        for source in item["representative_sources"]:
            url = str(source["url"]).strip()
            if url not in ledger_urls:
                raise ValueError(
                    f"dimension {number} representative source URL is not present in the source ledger: {url}"
                )


def validate_result_layer(data: dict[str, object]) -> dict[str, object]:
    layer = data.get("result_layer")
    if not isinstance(layer, dict):
        raise ValueError("result_layer is required for the +1 comprehensive result")
    if str(layer.get("name", "")).strip() != RESULT_LAYER_NAME:
        raise ValueError(f"result_layer name must be {RESULT_LAYER_NAME}")
    score = layer.get("composite_score")
    if score is not None and (
        isinstance(score, bool) or not isinstance(score, (int, float)) or not 0 <= score <= 100
    ):
        raise ValueError("result_layer composite_score must be null or a number from 0 to 100")
    for key in ("core_strengths", "core_risks", "bottom_line_issues"):
        value = layer.get(key)
        if not isinstance(value, list) or not value or any(not str(item).strip() for item in value):
            raise ValueError(f"result_layer {key} must be a non-empty text array")
    if not str(layer.get("evidence_sufficiency", "")).strip():
        raise ValueError("result_layer evidence_sufficiency is required")
    if str(layer.get("evaluation_confidence", "")).strip() not in EVIDENCE_STRENGTHS:
        raise ValueError("result_layer evaluation_confidence is invalid")
    if layer.get("bottom_line_override_checked") is not True:
        raise ValueError("result_layer must explicitly confirm the bottom-line non-compensation check")
    return layer


def add_dimension_text(document: object, item: object, style: str | None = None) -> None:
    payload = item if isinstance(item, dict) else {"text": str(item)}
    add_paragraph_with_links(document, payload, style=style)


def add_dimension_heading(document: object, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    font_size = 14 if level == 2 else 12
    for run in paragraph.runs:
        set_run_font(run, "Times New Roman", font_size, bold=True, east_asia="黑体")


def add_labeled_paragraph(document: object, label: str, value: str) -> object:
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, "黑体", 10.5, bold=True)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, "宋体", 10.5)
    return paragraph


def add_labeled_dimension_items(
    document: object,
    label: str,
    values: list[object],
) -> None:
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, "黑体", 10.5, bold=True)
    for index, value in enumerate(values):
        payload = value if isinstance(value, dict) else {"text": str(value), "links": []}
        if index:
            separator = paragraph.add_run("；")
            set_run_font(separator, "宋体", 10.5)
        text_run = paragraph.add_run(str(payload.get("text", "")))
        set_run_font(text_run, "宋体", 10.5)
        links = payload.get("links", [])
        if isinstance(links, list):
            for link in links:
                if not isinstance(link, dict):
                    continue
                url = str(link.get("url", "")).strip()
                if not re.match(r"^https?://", url):
                    continue
                spacer = paragraph.add_run(" ")
                set_run_font(spacer, "宋体", 10.5)
                add_hyperlink(paragraph, str(link.get("label", "来源")) or "来源", url)


def add_representative_sources(document: object, sources: list[dict[str, object]]) -> None:
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run("代表性来源：")
    set_run_font(label_run, "黑体", 10.5, bold=True)
    for index, source in enumerate(sources):
        if index:
            separator = paragraph.add_run("、")
            set_run_font(separator, "宋体", 10.5)
        add_hyperlink(paragraph, str(source["label"]), str(source["url"]))
    suffix = paragraph.add_run("。")
    set_run_font(suffix, "宋体", 10.5)


def add_compact_dimension_summary(
    document: object,
    item: dict[str, object],
) -> None:
    tendency = item.get("tendency_score")
    tendency_text = "数据不足" if tendency is None else f"{float(tendency):+.2f}"
    conversion = item.get("conversion_score")
    conversion_text = "数据不足" if conversion is None else f"{float(conversion):.1f}分"
    status_text = (
        "达到中高或高置信度优先目标"
        if item["evidence_status"] == "sufficient"
        else "达到明确选择的中置信度目标"
        if item["evidence_status"] == "sufficient_at_requested_target"
        else "达到中置信度并通过增强终止独立审计"
        if item["evidence_status"] == "sufficient_at_medium_after_audit"
        else "检索按受控条件终止，保留符合最低正式门槛的结果并披露缺口"
        if item["evidence_status"] == AUDIT_RETRIEVAL_TERMINATED_WITH_SHORTFALL
        else "经独立审计确认穷尽后仍低于最低要求"
    )
    rounds = item.get("targeted_deep_rounds", [])
    platform_status = {
        "scored": "已计分",
        "insufficient_platform_samples": "平台样本不足",
    }.get(str(item["platform_minimum_rule_status"]), str(item["platform_minimum_rule_status"]))
    paragraph = document.add_paragraph()
    for line_index, (label, value) in enumerate(
        (
            ("初始研究权重", f"{float(item['initial_research_weight']):.0%}"),
            ("跨平台等权倾向值", tendency_text),
            ("跨平台等权维度得分", conversion_text),
            ("证据充分度", str(item["evidence_strength"])),
            ("检索有效证据数", str(item["retrieved_evidence_units"])),
            ("主题覆盖证据数", str(item["topic_coverage_evidence_units"])),
            ("评分候选证据数", str(item["eligible_evidence_units"])),
            ("实际计分证据数", f"{item['scored_evidence_units']}（最低25）"),
            (
                "计分正中负构成",
                f"正面{item['scoring_positive_units']}、中性{item['scoring_neutral_units']}、负面{item['scoring_negative_units']}",
            ),
            (
                "有效计分平台",
                f"{item['dimension_scorable_platform_count']}/{item['run_included_platform_count']}；"
                f"候选平台{item['dimension_candidate_platform_count']}；{platform_status}",
            ),
            ("独立来源页数", f"{item['distinct_source_pages']}（最低12）"),
            ("来源类型数", f"{item['source_category_count']}（最低3）"),
            ("维度置信度", str(item["dimension_confidence"])),
            ("最低置信度要求", "中"),
            ("优先置信度目标", "中高/高"),
            ("维度证据状态", status_text),
            ("定向深检轮次", "、".join(str(value) for value in rounds) if rounds else "未触发或无需触发"),
            ("定向深检轮数", str(item["targeted_round_count"])),
            (
                "独立穷尽审计",
                "通过" if item.get("independent_exhaustion_audit_passed") is True else "不适用",
            ),
            (
                "中置信度终止审计",
                "通过" if item.get("medium_completion_audit_passed") is True else "不适用",
            ),
            ("中置信度终止依据", str(item.get("medium_completion_basis", "")) or "不适用"),
        )
    ):
        if line_index:
            breaker = paragraph.add_run()
            breaker.add_break()
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, "宋体", 10.5, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, "宋体", 10.5)
    breaker = paragraph.add_run()
    breaker.add_break()
    label_run = paragraph.add_run("代表性来源：")
    set_run_font(label_run, "宋体", 10.5, bold=True)
    for index, source in enumerate(item["representative_sources"]):
        if index:
            separator = paragraph.add_run("、")
            set_run_font(separator, "宋体", 10.5)
        add_hyperlink(paragraph, str(source["label"]), str(source["url"]))
    suffix = paragraph.add_run("。")
    set_run_font(suffix, "宋体", 10.5)


def add_dimension_analyses(document: object, analyses: list[dict[str, object]]) -> None:
    for number, item in enumerate(analyses, start=1):
        add_dimension_heading(document, f"5.{number} {item['name']}", level=2)
        add_labeled_paragraph(document, "维度定义", str(item["definition"]))
        add_dimension_heading(document, "主要正面评价", level=3)
        for paragraph in item["positive"]:
            add_dimension_text(document, paragraph)
        add_dimension_heading(document, "主要负面评价", level=3)
        for paragraph in item["negative"]:
            add_dimension_text(document, paragraph)
        add_dimension_heading(document, str(item["differences_title"]), level=3)
        for paragraph in item["differences"]:
            add_dimension_text(document, paragraph)
        for paragraph in item["fact_perception"]:
            add_dimension_text(document, paragraph)
        add_labeled_dimension_items(document, "作用机制", item["mechanism_analysis"])
        add_labeled_dimension_items(document, "反证与替代解释", item["counterevidence"])
        add_labeled_dimension_items(document, "不确定性与适用边界", item["uncertainty_boundary"])
        add_labeled_dimension_items(document, "综合判断", item["synthesis_judgment"])
        add_labeled_paragraph(document, "编码边界", str(item["boundary"]))
        for block in item.get("additional_analysis", []):
            if isinstance(block, dict):
                add_block(document, block)
        add_labeled_paragraph(document, "高频主题", "、".join(str(value) for value in item["themes"]))
        mention_text = (
            f"{item['mention_level']}，{item['mention_range']}。{item['mention_basis']}"
        )
        add_labeled_paragraph(document, "提及率等级", mention_text)
        add_compact_dimension_summary(document, item)


def validate_report_data(data: dict[str, object]) -> list[dict[str, object]]:
    title = str(data.get("title", "")).strip()
    if not title:
        raise ValueError("report title is required")
    if not str(data.get("subtitle", "")).strip():
        raise ValueError("report subtitle is required")
    if str(data.get("dimension_weight_scheme", "")).strip() != WEIGHT_SCHEME_NAME:
        raise ValueError(f"dimension_weight_scheme must be {WEIGHT_SCHEME_NAME}")
    declared_weights = data.get("dimension_weights")
    if not isinstance(declared_weights, dict) or set(declared_weights) != set(DIMENSION_WEIGHTS):
        raise ValueError("dimension_weights must contain exactly the seven canonical dimensions")
    for dimension, expected in DIMENSION_WEIGHTS.items():
        actual = declared_weights.get(dimension)
        if (
            isinstance(actual, bool)
            or not isinstance(actual, (int, float))
            or abs(float(actual) - expected) > 1e-12
        ):
            raise ValueError(f"dimension_weights/{dimension} must equal {expected:.0%}")
    retrieval_date = str(data.get("retrieval_date", "")).strip()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", retrieval_date):
        raise ValueError("retrieval_date must use YYYY-MM-DD")
    validate_research_scope(data)
    validate_research_run(data)
    result_layer = validate_result_layer(data)
    sections = data.get("sections")
    if not isinstance(sections, list) or len(sections) != len(SECTIONS):
        raise ValueError("report data must contain exactly 15 sections")
    normalized: list[dict[str, object]] = []
    for number, expected_title in enumerate(SECTIONS, start=1):
        item = sections[number - 1]
        if not isinstance(item, dict):
            raise ValueError(f"section {number} must be an object")
        if item.get("number") != number or str(item.get("title", "")).strip() != expected_title:
            raise ValueError(f"section {number} must be titled {expected_title}")
        blocks = item.get("blocks", [])
        if not isinstance(blocks, list):
            raise ValueError(f"section {number} blocks must be an array")
        if number == 5 and any(
            isinstance(block, dict) and block.get("type") not in {"paragraph", "quote", "bullets"}
            for block in blocks
        ):
            raise ValueError(
                "section 5 blocks may contain introductory narrative only; use dimension_analyses for all seven dimensions"
            )
        if number == 14 and any(
            not isinstance(block, dict) or block.get("type") not in {"paragraph", "quote"}
            for block in blocks
        ):
            raise ValueError("section 14 blocks may contain source-list guidance paragraphs only")
        if number == 15 and blocks:
            raise ValueError("section 15 blocks must be empty; use machine_readable_line")
        if number not in {5, 14, 15}:
            validate_section_contract(number, blocks)
        normalized.append(item)
    summary_table = next(
        block
        for block in normalized[5]["blocks"]
        if isinstance(block, dict) and block.get("type") == "table"
    )
    summary_rows = summary_table.get("rows", [])
    assert isinstance(summary_rows, list)
    displayed_weights: dict[str, float] = {}
    for row in summary_rows:
        if not isinstance(row, list) or len(row) < 2:
            raise ValueError("section 6 dimension summary rows must include dimension and initial research weight")
        dimension = str(row[0]).strip()
        raw_weight = row[1]
        text_weight = str(raw_weight).strip()
        try:
            displayed_weights[dimension] = (
                float(text_weight[:-1]) / 100 if text_weight.endswith("%") else float(raw_weight)
            )
        except (TypeError, ValueError):
            raise ValueError(f"section 6 initial research weight for {dimension} is not numeric") from None
    if set(displayed_weights) != set(DIMENSION_WEIGHTS):
        raise ValueError("section 6 dimension summary must contain exactly the seven canonical dimensions")
    for dimension, expected in DIMENSION_WEIGHTS.items():
        if abs(displayed_weights[dimension] - expected) > 1e-12:
            raise ValueError(
                f"section 6 initial research weight for {dimension} must equal {expected:.0%}"
            )
    result_block = next(
        block
        for block in normalized[11]["blocks"]
        if isinstance(block, dict) and block.get("type") == "judgments"
    )
    result_items = result_block["items"]
    result_values = {
        str(item["label"]): str(item["value"]).strip()
        for item in result_items
        if isinstance(item, dict)
    }
    expected_result_values = {
        "综合分": "数据不足" if result_layer["composite_score"] is None else str(result_layer["composite_score"]),
        "核心优势": "；".join(str(value) for value in result_layer["core_strengths"]),
        "核心风险": "；".join(str(value) for value in result_layer["core_risks"]),
        "底线性问题": "；".join(str(value) for value in result_layer["bottom_line_issues"]),
        "证据充分度": str(result_layer["evidence_sufficiency"]),
        "评价置信度": str(result_layer["evaluation_confidence"]),
    }
    if result_values != expected_result_values:
        raise ValueError("section 12 +1 result fields must exactly match the structured result_layer")
    dimensions = validate_dimension_analyses(data)
    dimension_section_text = "".join(
        str(item["name"])
        + text_items_text(item["positive"])
        + text_items_text(item["negative"])
        + str(item["differences_title"])
        + text_items_text(item["differences"])
        + text_items_text(item["fact_perception"])
        + text_items_text(item["mechanism_analysis"])
        + text_items_text(item["counterevidence"])
        + text_items_text(item["uncertainty_boundary"])
        + text_items_text(item["synthesis_judgment"])
        + "".join(str(value) for value in item["themes"])
        + str(item["mention_level"])
        + str(item["mention_range"])
        + str(item["mention_basis"])
        + str(item["evidence_strength"])
        for item in dimensions
    )
    if cjk_count(dimension_section_text) < 5000:
        raise ValueError("section 5 detailed analysis must contain at least 5000 Chinese characters")
    analytical_records: list[tuple[str, str]] = []
    for section in normalized[:13]:
        number = int(section["number"])
        blocks = section.get("blocks", [])
        if not isinstance(blocks, list):
            continue
        for block_index, block in enumerate(blocks, start=1):
            if not isinstance(block, dict) or block.get("type") not in {"paragraph", "quote", "bullets", "numbered"}:
                continue
            analytical_records.append(
                (f"section {number} block {block_index}", block_text(block))
            )
    for dimension_index, item in enumerate(dimensions, start=1):
        for field in (
            "positive", "negative", "differences", "fact_perception", "mechanism_analysis",
            "counterevidence", "uncertainty_boundary", "synthesis_judgment",
        ):
            analytical_records.append(
                (f"dimension {dimension_index} {field}", text_items_text(item[field]))
            )
    assert_no_near_duplicate_narratives(
        analytical_records,
        context="report analytical prose",
    )
    validate_english_gloss_contract(data)
    report_table_count = sum(
        1
        for section in normalized
        for block in section.get("blocks", [])
        if isinstance(block, dict) and block.get("type") == "table"
    )
    if report_table_count != 11:
        raise ValueError("sections 1 to 13 must contain exactly eleven analytical tables before the automatic source table")
    machine_line = str(data.get("machine_readable_line", "")).strip()
    if len(machine_line.split("｜")) != 19:
        raise ValueError("machine_readable_line must contain exactly 19 fields separated by ｜")
    return normalized


def chinese_date(value: str) -> str:
    year, month, day = value.split("-")
    return f"{int(year)}年{int(month)}月{int(day)}日"


def validate_report_score_truth(
    analyses: list[dict[str, object]],
    research_run: dict[str, object],
    scores: dict[str, object],
) -> None:
    if scores.get("status") != "valid":
        raise ValueError("formal DOCX requires a valid composite-score truth file")
    scoring_input = scores.get("formal_scoring_input")
    if not isinstance(scoring_input, dict):
        raise ValueError("composite-score truth lacks its formal scoring gate input")
    verify_scoring_input_gate(scoring_input)
    if scores.get("task_run_id") != research_run.get("task_run_id"):
        raise ValueError("report and composite-score truth use different task_run_id values")
    cross = scores.get("cross_platform")
    dimensions = cross.get("dimensions") if isinstance(cross, dict) else None
    if not isinstance(dimensions, dict):
        raise ValueError("composite-score truth lacks cross-platform dimension results")
    for item in analyses:
        name = str(item["name"])
        score_item = dimensions.get(name)
        if not isinstance(score_item, dict):
            raise ValueError(f"composite-score truth lacks dimension: {name}")
        formal_score = score_item.get("platform_equal_score")
        expected_tendency = None if formal_score is None else float(formal_score) / 10 - 5
        for field, expected in (
            ("conversion_score", formal_score),
            ("tendency_score", expected_tendency),
        ):
            actual = item.get(field)
            if actual is None or expected is None:
                if actual is not expected:
                    raise ValueError(f"report {name} {field} differs from the formal score truth")
            elif abs(float(actual) - float(expected)) > 1e-9:
                raise ValueError(f"report {name} {field} differs from the formal score truth")


def build_docx(
    report_data_path: Path,
    sources_path: Path,
    scores_path: Path,
    output_path: Path,
) -> dict[str, object]:
    data = json.loads(report_data_path.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict):
        raise ValueError("report data root must be an object")
    from input_safety import assert_public_text
    assert_public_text(data)
    sections = validate_report_data(data)
    dimension_analyses = validate_dimension_analyses(data)
    research_run = validate_research_run(data)
    scores = json.loads(scores_path.read_text(encoding="utf-8-sig"))
    if not isinstance(scores, dict):
        raise ValueError("composite-score truth root must be an object")
    validate_report_score_truth(dimension_analyses, research_run, scores)
    sources = read_csv(sources_path)
    assert_public_text(sources)
    from host_receipts import proof_errors
    proof_issues = [issue for source in sources for issue in proof_errors(source)]
    if proof_issues:
        raise ValueError("report source host proof failed: " + "; ".join(proof_issues))
    source_run_ids = {source.get("task_run_id", "") for source in sources if source.get("task_run_id", "")}
    if source_run_ids != {str(research_run["task_run_id"])}:
        raise ValueError(
            "source ledger task_run_id must exactly match research_run task_run_id"
        )
    validate_representative_sources_in_ledger(dimension_analyses, sources)
    document = Document()
    configure_document(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(str(data["title"]))
    set_run_font(run, "Times New Roman", 24, bold=True, east_asia="黑体")
    subtitle_text = str(data.get("subtitle", "")).strip()
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(subtitle_text)
    set_run_font(run, "Times New Roman", 14, east_asia="黑体")
    date_paragraph = document.add_paragraph(style="Date")
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_paragraph.add_run(f"检索日期：{chinese_date(str(data['retrieval_date']))}")
    set_run_font(date_run, "Times New Roman", 11, east_asia="宋体")
    document.add_page_break()
    toc_heading = document.add_paragraph("目录", style="TOC Heading")
    for run in toc_heading.runs:
        set_run_font(run, "Times New Roman", 16, bold=True, east_asia="黑体")
    toc_items = ["研究口径说明"] + [
        f"{index}. {title}" for index, title in enumerate(SECTIONS, start=1)
    ]
    for item in toc_items:
        paragraph = document.add_paragraph(item)
        for run in paragraph.runs:
            set_run_font(run, "宋体", 10.5)
    document.add_page_break()
    scope_heading = document.add_heading("研究口径说明", level=1)
    for run in scope_heading.runs:
        set_run_font(run, "Times New Roman", 16, bold=True, east_asia="黑体")
    for item in data["research_scope"]:
        add_dimension_text(document, item)
    add_labeled_paragraph(document, "网络采集证明等级",
        f"正式来源台账中，宿主签章核验来源{sum(s.get('network_proof_level') == 'host_verified' for s in sources)}条，"
        f"工具可追溯来源{sum(s.get('network_proof_level') == 'tool_traceable' for s in sources)}条。"
        "两类均已复核本次任务、页面读取记录、响应快照及证据定位；工具可追溯不表示已取得宿主密码学签章或连接地址证明。"
        "无本次实际工具读取依据的材料不纳入正式证据、评分、来源多样性、穷尽或受限交付依据。")
    minimum_effective_samples = int(research_run["minimum_effective_samples"])
    actual_effective_samples = int(research_run["actual_effective_samples"])
    if actual_effective_samples < minimum_effective_samples:
        sample_gap = minimum_effective_samples - actual_effective_samples
        deep_rounds = int(research_run["deep_iteration_rounds"])
        sample_status = str(research_run["sample_target_status"])
        terminal_text = (
            "检索已达到机器审计确认的受控终止条件，本报告仅形成受限交付，"
            "未将未达到的目标解释为已经满足。"
            if sample_status == "retrieval_terminated_with_shortfall"
            else "检索经机器审计确认已经穷尽，本报告保留样本不足结论。"
        )
        add_labeled_paragraph(
            document,
            "样本目标与检索终态",
            (
                f"本次共完成{deep_rounds}轮深度迭代；实际计分证据为"
                f"{actual_effective_samples}个，目标为{minimum_effective_samples}个，"
                f"样本缺口为{sample_gap}个。{terminal_text}"
            ),
        )
    semantic = research_run["semantic_quantification"]
    add_labeled_paragraph(
        document,
        "锁定评价协议",
        (
            f"评价协议标识：{semantic['evaluation_protocol_version']}；"
            f"语义代码簿标识：{semantic['semantic_codebook_version']}；"
            f"SHA-256（安全散列算法校验值）：{semantic['semantic_codebook_sha256']}；"
            f"协议发布日期：{semantic['protocol_released_at']}。"
            "规则在Skill（技能）版本修订期完成发布校准，本次检索任务不在线更新或改写规则。"
        ),
    )
    for index, section in enumerate(sections, start=1):
        heading = document.add_heading(f"{index}. {SECTIONS[index - 1]}", level=1)
        for run in heading.runs:
            set_run_font(run, "Times New Roman", 16, bold=True, east_asia="黑体")
        if index == 5:
            for block in section.get("blocks", []):
                if isinstance(block, dict):
                    add_block(document, block)
            add_dimension_analyses(document, dimension_analyses)
        elif index == 14:
            for block in section.get("blocks", []):
                if isinstance(block, dict) and block.get("type") != "table":
                    add_block(document, block)
            add_source_table(document, sources)
        elif index == 15:
            line = document.add_paragraph(str(data["machine_readable_line"]))
            for run in line.runs:
                set_run_font(run, "Times New Roman", 9)
        else:
            for block in section.get("blocks", []):
                if isinstance(block, dict):
                    add_block(document, block)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {
        "status": "created",
        "output": str(output_path.resolve()),
        "section_count": 15,
        "source_rows": len(sources),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Structured report JSON")
    parser.add_argument("--sources", required=True, type=Path)
    parser.add_argument("--scores", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--state", required=True, type=Path)
    parser.add_argument("--writer-ledger", required=True, type=Path)
    parser.add_argument("--truth-freeze", required=True, type=Path)
    parser.add_argument("--preflight", required=True, type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        authorize_runtime_write(
            state_path=args.state,
            writer_script_id="build_docx_report.py",
            output_role="docx_report",
            expected_phase="REPORT_BUILD",
            output_path=args.output,
        )
        freeze = verify_truth_freeze(
            state_path=args.state,
            writer_ledger_path=args.writer_ledger,
            manifest_path=args.truth_freeze,
            require_scoring=True,
        )
        assert_frozen_artifact_path(freeze, role="source_ledger", path=args.sources)
        assert_frozen_artifact_path(freeze, role="scoring_output", path=args.scores)
        verify_preflight(
            state_path=args.state,
            writer_ledger_path=args.writer_ledger,
            preflight_path=args.preflight,
        )
        verify_artifact_writer(
            state_path=args.state,
            writer_ledger_path=args.writer_ledger,
            output_role="report_data",
            output_path=args.input,
        )
        result = build_docx(args.input, args.sources, args.scores, args.output)
        register_protected_artifact(
            state_path=args.state,
            writer_ledger_path=args.writer_ledger,
            writer_script_id="build_docx_report.py",
            output_role="docx_report",
            output_path=args.output,
            input_paths={
                "report_data": args.input,
                "sources": args.sources,
                "scores": args.scores,
                "truth_freeze": args.truth_freeze,
                "preflight": args.preflight,
            },
            expected_phase="REPORT_BUILD",
        )
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(json.dumps({"status": "invalid", "errors": [str(exc)]}, ensure_ascii=False, indent=2))
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
